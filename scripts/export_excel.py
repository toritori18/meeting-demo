"""DB から「画面一覧・機能一覧.xlsx」を生成する。

sys-staff が中身（どの画面・どの機能・どの項目か）を決めてDBに登録し、
このスクリプトが体裁を作る。エージェントに xlsx バイナリは書けないための分担。

    python scripts/export_excel.py
    python scripts/export_excel.py --open   # 生成後に Excel で開く
"""
from __future__ import annotations

import argparse
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db_init import DOCS, connect, now

OUT = DOCS / "画面一覧・機能一覧.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
HEAD_FONT = Font(bold=True)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)

STATUS_FILL = {
    "承認済": PatternFill("solid", fgColor="E2EFDA"),
    "差戻し": PatternFill("solid", fgColor="FCE4E4"),
    "未着手": PatternFill("solid", fgColor="F2F2F2"),
}
JUDGMENT_FILL = {
    "採用": PatternFill("solid", fgColor="E2EFDA"),
    "対象外": PatternFill("solid", fgColor="F2F2F2"),
    "次期": PatternFill("solid", fgColor="FFF2CC"),
    "保留": PatternFill("solid", fgColor="FCE4E4"),
}


def _meeting_label(con, mid):
    if not mid:
        return ""
    row = con.execute("SELECT seq FROM meeting WHERE id=?", (mid,)).fetchone()
    return "第{}回".format(row["seq"]) if row else ""


def _sheet(wb, title, headers, widths, rows, fills=None, fill_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", horizontal="center")
    for row in rows:
        ws.append(["" if v is None else v for v in row])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.border = BORDER
            cell.alignment = WRAP
        if fills and fill_col is not None:
            key = r[fill_col].value
            if key in fills:
                r[fill_col].fill = fills[key]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build(con) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    screens = con.execute("SELECT * FROM screen ORDER BY code").fetchall()
    rows = []
    for s in screens:
        n_open = con.execute(
            "SELECT COUNT(*) AS n FROM screen_review WHERE screen_id=? AND resolved=0",
            (s["id"],)).fetchone()["n"]
        items = con.execute(
            "SELECT name FROM screen_item WHERE screen_id=? ORDER BY seq",
            (s["id"],)).fetchall()
        rows.append([
            s["code"], s["name"], s["area"], s["summary"],
            "、".join(i["name"] for i in items),
            s["designer"], s["status"],
            _meeting_label(con, s["approved_meeting_id"]),
            s["approved_at"] or "", n_open, s["note"],
        ])
    _sheet(wb, "画面一覧",
           ["画面ID", "画面名", "機能区分", "概要", "主な項目", "設計者",
            "承認状態", "承認会議", "承認日", "未解消指摘", "備考"],
           [10, 22, 12, 40, 34, 11, 15, 10, 18, 11, 26],
           rows, STATUS_FILL, 6)

    features = con.execute("SELECT * FROM feature ORDER BY code").fetchall()
    _sheet(wb, "機能一覧",
           ["機能ID", "機能名", "機能区分", "対応画面", "概要", "優先度",
            "判定", "判定会議", "判定日", "備考"],
           [10, 24, 12, 20, 44, 9, 9, 10, 18, 26],
           [[f["code"], f["name"], f["area"], f["screens"], f["summary"],
             f["priority"], f["judgment"], _meeting_label(con, f["judged_meeting_id"]),
             f["judged_at"] or "", f["note"]] for f in features],
           JUDGMENT_FILL, 6)

    item_rows = []
    for s in screens:
        if s["status"] != "承認済":
            continue  # 承認ゲート: 未承認の画面は項目定義を出さない
        for it in con.execute(
                "SELECT * FROM screen_item WHERE screen_id=? ORDER BY seq",
                (s["id"],)).fetchall():
            item_rows.append([
                s["code"], s["name"], it["seq"], it["name"], it["datatype"],
                it["length"], it["required"], it["default_value"],
                it["validation"], it["note"]])
    ws = _sheet(wb, "画面項目定義",
                ["画面ID", "画面名", "No", "項目名", "型", "桁", "必須",
                 "初期値", "バリデーション", "備考"],
                [10, 20, 6, 22, 12, 8, 8, 16, 34, 26],
                item_rows)
    if not item_rows:
        ws["A2"] = "承認済みの画面がまだないため空です（承認ゲート）"
        ws.merge_cells("A2:J2")

    revs = con.execute(
        "SELECT r.*, m.seq AS mseq FROM revision r "
        "LEFT JOIN meeting m ON m.id = r.meeting_id ORDER BY r.id").fetchall()
    if revs:
        rev_rows = [[r["version"], r["dated"],
                     "第{}回".format(r["mseq"]) if r["mseq"] else "", r["body"]]
                    for r in revs]
    else:
        rev_rows = [["0.1", now()[:10], "", "初版（自動生成）"]]
    _sheet(wb, "改訂履歴", ["版", "日付", "会議", "変更内容"],
           [8, 14, 10, 80], rev_rows)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)


def main() -> None:
    ap = argparse.ArgumentParser(description="画面一覧・機能一覧のExcelを生成")
    ap.add_argument("--open", action="store_true", help="生成後にExcelで開く")
    args = ap.parse_args()

    con = connect()
    try:
        build(con)
    finally:
        con.close()

    print("生成: {}".format(OUT))
    if args.open:
        try:
            if sys.platform == "win32":
                os.startfile(OUT)  # noqa: S606
            else:
                subprocess.run(["xdg-open", str(OUT)], check=False)
        except Exception as exc:
            print("開けませんでした: {}".format(exc))


if __name__ == "__main__":
    main()
